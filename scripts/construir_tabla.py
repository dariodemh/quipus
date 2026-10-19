"""Aplana los 651 xlsx de KFG en analisis/datos/{cuerdas,quipus,grupos}.csv."""
import csv, re, sys
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parents[1]
XLSX = RAIZ / "data/external/data-science-khipu-article/data/KFG_DB_Excel"
OUT = RAIZ / "analisis/datos"
GRUPO = re.compile(r"([\d.]+)cm (?:group of (\d+) \w+ \((\d+)-(\d+)\)(?: space of ([\d.]+)cm)?|1 \((\d+)\))")

def kv(ws):
    d = {}
    for (v, *_) in ws.iter_rows(values_only=True):
        if isinstance(v, str) and ":" in v and not v.startswith("!"):
            k, _, x = v.partition(":"); d[k.strip()] = x.strip()
    return d

def main():
    OUT.mkdir(parents=True, exist_ok=True)
    cuerdas, quipus, grupos = [], [], []
    for f in sorted(XLSX.glob("*.xlsx")):
        wb = openpyxl.load_workbook(f, read_only=True)
        k = f.stem
        meta = kv(wb["Khipu"]); pc = kv(wb["PrimaryCord"])
        quipus.append({"kfg": k, "okr": meta.get("OKR_Name", ""), "museo": meta.get("Museum Name", ""),
                       "procedencia": meta.get("Provenance", ""), "region": meta.get("Region", ""),
                       "pc_largo": pc.get("Length", ""), "pc_color": pc.get("Color", ""), "pc_estructura": pc.get("Structure", "")})
        # grupos de colgantes: número de cuerda de primer nivel -> índice de grupo
        idx = {}
        for g, (v, *_) in enumerate(r for r in wb["Clusters"].iter_rows(values_only=True) if r[0] and not str(r[0]).startswith("!")):
            m = GRUPO.match(str(v))
            if not m: continue
            a, b = (int(m[3]), int(m[4])) if m[2] else (int(m[6]), int(m[6]))
            grupos.append({"kfg": k, "grupo": g, "pos_cm": m[1], "desde": a, "hasta": b, "espacio_cm": m[5] or ""})
            for n in range(a, b + 1): idx[n] = g
        filas = list(wb["Cords"].iter_rows(values_only=True))
        cab = filas[0]
        for orden, r in enumerate(filas[1:]):
            d = dict(zip(cab, r))
            nom = str(d["Cord_Name"])
            m = re.match(r"([a-z]+)(\d+)", nom)
            raiz = int(m[2]) if m else None
            padre = nom[: max(nom.rfind("s"), 0)] if nom.count("s") and re.search(r"\ds\d", nom) else ""
            cuerdas.append({"kfg": k, "orden": orden, "cuerda": nom, "tipo": m[1] if m else "",
                            "nivel": len(re.findall(r"s\d+", nom)), "padre": padre,
                            "grupo": idx.get(raiz, "") if m and m[1] == "p" else "",
                            "valor": d["Value"], "alt": d["Alt_Value"] or "", "nudos": d["Knots"] or "",
                            "color": d["Color"] or "", "torsion": d["Twist"] or "", "union": d["Attachment"] or "",
                            "largo": d["Length"] or "", "terminacion": d["Termination"] or "", "notas": (d["Notes"] or "").replace("\n", " ")})
    for nombre, filas in (("cuerdas", cuerdas), ("quipus", quipus), ("grupos", grupos)):
        with open(OUT / f"{nombre}.csv", "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=list(filas[0])); w.writeheader(); w.writerows(filas)
        print(nombre, len(filas))

if __name__ == "__main__":
    main()
